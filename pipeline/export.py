"""Excel export following Herrenknecht's BoreholeDataFormat.xlsx.

The format is a field list in three groups (metadata, log data, test results). Log data mixes
per-layer fields (lithology, stratigraphy) with depth-indexed ones (samples, core runs, SPT,
groundwater), so the workbook has one sheet per kind of row, all keyed by BoreholeID. Column names
and units are the ones of the format; columns the format does not have (WGS84 position, review
state, source pages) come last.
"""

from __future__ import annotations

import io
from collections import defaultdict
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from pyproj import CRS
from pyproj.exceptions import CRSError

from schema.borehole import DIN4023_LABEL_DE, Borehole, LabParameter, LabResult

# Test result columns of the format, in its order (names as in the format, units in brackets).
TEST_COLUMNS: list[tuple[LabParameter, str]] = [
    (LabParameter.water_content, "Watercontent [%]"),
    (LabParameter.consistency, "Cosistency [%]"),
    (LabParameter.unit_weight, "Unitweight [kN/m³]"),
    (LabParameter.ucs, "UCS [MPa]"),
    (LabParameter.tensile_strength, "Tensile strength [MPa]"),
    (LabParameter.cai, "CAI [-]"),
    (LabParameter.rmr, "RMR [-]"),
    (LabParameter.q, "Q [-]"),
    (LabParameter.gsi, "GSI [-]"),
    (LabParameter.friction_angle, "FrictionAngle [°]"),
    (LabParameter.cohesion, "Cohesion [kPa]"),
    (LabParameter.fines, "Grainsize-fines [%]"),
    (LabParameter.liquid_limit, "LiquidLimit [%]"),
    (LabParameter.plastic_limit, "PlasticityLimit [%]"),
    (LabParameter.plasticity_index, "PlasticityIndex [-]"),
    (LabParameter.consistency_index, "ConsistencyIndex [-]"),
    (LabParameter.hydraulic_conductivity, "HydraulicCoefficient [m/s]"),
    (LabParameter.e_modulus, "E-Modulus [MPa]"),
    (LabParameter.poissons_ratio, "PoissonsRatio [-]"),
    # not in the format, kept so no extracted value is lost
    (LabParameter.dry_unit_weight, "DryUnitweight [kN/m³]"),
    (LabParameter.undrained_shear_strength, "UndrainedShearStrength [kPa]"),
    (LabParameter.sand, "Grainsize-sand [%]"),
    (LabParameter.gravel, "Grainsize-gravel [%]"),
    (LabParameter.point_load_index, "PointLoadIndex Is50 [MPa]"),
    (LabParameter.other, "Other (see raw)"),
]

FIELDS = [  # (sheet, column, description) for the 'Fields' sheet
    ("Metadata", "X-Coordinate / Y-Coordinate [m/°]", "As printed, in CoordinateSystem (easting/northing, or lon/lat)"),
    ("Metadata", "CoordinateSystem", "Resolved EPSG code and name; 'as printed' when it could not be resolved"),
    ("Metadata", "Surface Elevation [mASL]", "Ground level as printed; see ElevationDatum"),
    ("Metadata", "Latitude/Longitude WGS84 [°]", "Position transformed to WGS84 (not in the format)"),
    ("Metadata", "ReviewStatus / Score", "Human-in-the-loop state and release score 0-100 (not in the format)"),
    ("Logdata", "Lithology/SoilName", "Description as printed (original language)"),
    ("Logdata", "Soil/Rockprofile", "DIN 4023 class of the main constituent (pattern key) and material colour"),
    ("Logdata", "Core", "yes if a core run covers the layer, no if the hole has core runs elsewhere, empty if unknown"),
    ("TestResults", "Value columns", "Converted to the unit in the header; the raw value is in 'Raw values'"),
]


@lru_cache(maxsize=None)
def crs_name(epsg: int) -> str:
    try:
        return f"EPSG:{epsg} ({CRS.from_epsg(epsg).name})"
    except CRSError:
        return f"EPSG:{epsg}"


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            if isinstance(value, str):
                # Text from documents and reviewers: a control character would fail the whole workbook, and a
                # leading '=' would make openpyxl store it as a formula.
                ws.cell(r, c, ILLEGAL_CHARACTERS_RE.sub("", value)).data_type = "s"
            else:
                ws.cell(r, c, value)
    ws.freeze_panes = "A2"
    for i, name in enumerate(header, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(12, len(name) + 2), 60)


def _core(bh: Borehole, top: float, bottom: float) -> str | None:
    if not bh.core_runs:
        return None
    covered = any(c.top_m <= top + 0.01 and bottom <= c.bottom_m + 0.01 for c in bh.core_runs)
    return "yes" if covered else "no"


def workbook(records: list[tuple[Borehole, str | None]]) -> bytes:
    """records: (borehole, source file) pairs. Returns the .xlsx file content."""
    wb = Workbook()
    wb.remove(wb.active)
    meta, logs, samples, runs, spts, water, tests = [], [], [], [], [], [], []
    for bh, source_file in records:
        loc = bh.location
        crs = crs_name(loc.crs_epsg) if loc.crs_epsg else loc.crs_proj or (f"as printed: {loc.crs_raw}" if loc.crs_raw else None)
        pages = sorted({s.page for s in bh.sources if s.page is not None and (not source_file or s.file == source_file)})
        meta.append([
            bh.borehole_id, bh.borehole_name, bh.location_name, bh.start_date, bh.end_date, bh.project_id,
            bh.project_name, bh.surface_elevation_m, loc.easting, loc.northing, crs, bh.depth_m, bh.length_m,
            bh.notes, loc.coords_raw, bh.elevation_datum, loc.lat, loc.lon, loc.method.value if loc.method else None,
            bh.review_status.value, bh.reviewed_by, round((bh.score or 0) * 100), source_file,
            ";".join(map(str, pages)) or None,
        ])
        for layer in bh.layers:
            logs.append([
                bh.borehole_id, layer.top_m, layer.bottom_m, layer.stratigraphic_name, layer.engineering_unit,
                layer.geologic_age, layer.description_original,
                f"{layer.primary.value} ({DIN4023_LABEL_DE.get(layer.primary, '')})" + (f", {layer.color}" if layer.color else ""),
                _core(bh, layer.top_m, layer.bottom_m), layer.stratigraphic_code, layer.description_en,
                layer.source.page if layer.source else None,
            ])
        samples += [[bh.borehole_id, s.name, s.top_m, s.bottom_m, s.type.value if s.type else s.type_raw]
                    for s in bh.samples]
        runs += [[bh.borehole_id, c.top_m, c.bottom_m, c.total_core_recovery_pct, c.rqd_pct, c.solid_core_recovery_pct]
                 for c in bh.core_runs]
        spts += [[bh.borehole_id, s.depth_m, "R" if s.refusal and s.n_value is None else s.n_value, s.raw] for s in bh.spt]
        water += [[bh.borehole_id, g.depth_m, g.kind, g.observed_on, "dry" if g.dry else None, g.reference, g.raw]
                  for g in bh.groundwater]
        by_sample: dict[tuple, list[dict]] = defaultdict(list)
        for r in bh.lab_results:
            rows = by_sample[(r.sample_name, r.top_m, r.bottom_m)]
            slot = next((values for values in rows if r.parameter not in values), None)
            if slot is None:
                slot = {}
                rows.append(slot)
            slot[r.parameter] = r
        for (name, top, bottom), rows in by_sample.items():
            for values in rows:
                raw = "; ".join(f"{r.parameter_raw or r.parameter.value}: {r.value_raw}" for r in values.values())
                tests.append([bh.borehole_id, name, top, bottom,
                              *[_test_value(values[p]) if p in values else None for p, _ in TEST_COLUMNS], raw])

    _sheet(wb, "Metadata", [
        "BoreholeID", "BoreholeName", "LocationName", "StartDate", "EndDate", "ProjectID", "ProjectName",
        "Surface Elevation [mASL]", "X-Coordinate [m/°]", "Y-Coordinate [m/°]", "CoordinateSystem",
        "BoreholeDepth [m]", "BoreholeLength [m]", "Notes", "Coordinates as printed", "ElevationDatum",
        "Latitude WGS84 [°]", "Longitude WGS84 [°]", "PositionMethod", "ReviewStatus", "ReviewedBy", "Score",
        "SourceFile", "SourcePages",
    ], meta)
    _sheet(wb, "Logdata", [
        "BoreholeID", "Top [m]", "Bottom [m]", "StratigraphicName", "EngineeringSoilUnit", "GeologicAge",
        "Lithology/SoilName", "Soil/Rockprofile", "Core [yes/no]", "StratigraphicCode", "Description (English)",
        "SourcePage",
    ], logs)
    _sheet(wb, "Samples", ["BoreholeID", "SampleName", "SampleDepth [m]", "SampleBottom [m]", "SampleType"], samples)
    _sheet(wb, "CoreRuns", ["BoreholeID", "Top [m]", "Bottom [m]", "TotalCoreRecovery [%]", "RQD [%]",
                            "SolidCoreRecovery [%]"], runs)
    _sheet(wb, "SPT", ["BoreholeID", "Depth [m]", "SPT [n]", "As printed"], spts)
    _sheet(wb, "Groundwater", ["BoreholeID", "Groundwaterlevel [m]", "Kind", "Date", "Dry", "Reference",
                               "As printed"], water)
    _sheet(wb, "TestResults", ["BoreholeID", "SampleName", "SampleDepth [m]", "SampleBottom [m]",
                               *[name for _, name in TEST_COLUMNS], "Raw values"], tests)
    _sheet(wb, "Fields", ["Sheet", "Column", "Meaning"], [list(f) for f in FIELDS])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _test_value(result: LabResult) -> float | str:
    if result.qualifier == "range":
        return f"{result.value:g}–{result.value_upper:g}" if result.value_upper is not None else result.value_raw or "range"
    return f"{result.qualifier}{result.value:g}" if result.qualifier else result.value
