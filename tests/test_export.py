"""Excel export in the BoreholeDataFormat layout (synthetic data)."""

import io

from openpyxl import load_workbook

from pipeline.export import workbook
from schema.borehole import (SPT, Borehole, CoreRun, GeorefMethod, GroundwaterObservation, LabResult, Layer,
                             Location, Sample, SourceRef)


def test_workbook_sheets_and_values():
    src = SourceRef(file="demo.pdf", page=2)
    bh = Borehole(
        borehole_id="demo:BH1", borehole_name="BH 1", depth_m=6.0, surface_elevation_m=35.7,
        location=Location(coords_raw="E 334000 N 6252000", easting=334000.0, northing=6252000.0, crs_epsg=28356,
                          lat=-33.85, lon=151.21, method=GeorefMethod.stated_coordinates),
        layers=[Layer(top_m=0, bottom_m=2, description_original="Clayey SAND", primary="sand", source=src),
                Layer(top_m=2, bottom_m=6, description_original="Sandstone", primary="sandstone", source=src)],
        core_runs=[CoreRun(top_m=2, bottom_m=6, total_core_recovery_pct=95, rqd_pct=60)],
        samples=[Sample(name="U1", top_m=1.0, bottom_m=1.45)],
        spt=[SPT(depth_m=1.0, n_value=12, raw="4,5,7")],
        groundwater=[GroundwaterObservation(depth_m=1.8, kind="standing")],
        lab_results=[LabResult(parameter="ucs", value=24.5, value_raw="24,5 MPa", top_m=3.0, sample_name="C2"),
                     LabResult(parameter="cai", value=2.1, value_raw="2.1", top_m=3.0, sample_name="C2")],
        sources=[src],
    )
    wb = load_workbook(io.BytesIO(workbook([(bh, "demo.pdf")])))
    assert wb.sheetnames == ["Metadata", "Logdata", "Samples", "CoreRuns", "SPT", "Groundwater", "TestResults", "Fields"]
    meta = {h.value: c.value for h, c in zip(wb["Metadata"][1], wb["Metadata"][2])}
    assert meta["X-Coordinate [m/°]"] == 334000 and meta["CoordinateSystem"].startswith("EPSG:28356")
    assert meta["SourcePages"] == "2"
    logs = [[c.value for c in row] for row in wb["Logdata"].iter_rows(min_row=2)]
    assert [r[8] for r in logs] == ["no", "yes"]  # only the rock layer was cored
    tests = {h.value: c.value for h, c in zip(wb["TestResults"][1], wb["TestResults"][2])}
    assert tests["UCS [MPa]"] == 24.5 and tests["CAI [-]"] == 2.1 and tests["SampleName"] == "C2"


def test_workbook_preserves_replicates_other_parameters_and_qualifiers():
    results = [LabResult(parameter="other", parameter_raw="Analyte A", value=1, value_raw="1 mg/l", top_m=2),
               LabResult(parameter="other", parameter_raw="Analyte B", value=2, value_raw="2 mg/l", top_m=2),
               LabResult(parameter="ucs", value=5, qualifier="<", value_raw="<5 MPa", top_m=2),
               LabResult(parameter="ucs", value=6, value_upper=9, qualifier="range", value_raw="6–9 MPa", top_m=2)]
    bh = Borehole(borehole_id="r:BH1", borehole_name="BH1", lab_results=results)
    wb = load_workbook(io.BytesIO(workbook([(bh, "r.pdf")])))
    rows = [dict(zip([c.value for c in wb["TestResults"][1]], row)) for row in
            wb["TestResults"].iter_rows(min_row=2, values_only=True)]
    assert len(rows) == 2
    assert {r["UCS [MPa]"] for r in rows} == {"<5", "6–9"}
    assert "Analyte A" in rows[0]["Raw values"] and "Analyte B" in rows[1]["Raw values"]
