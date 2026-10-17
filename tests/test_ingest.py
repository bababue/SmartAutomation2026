"""Ingestion of DOCX and plain-text logs (synthetic files)."""

import shutil

import docx
import pymupdf
import pytest

from pipeline import config, ingest


def _pdf(path, *texts):
    pdf = pymupdf.open()
    for text in texts:
        pdf.new_page().insert_text((72, 72), text)
    pdf.save(path)


def test_docx_uses_converted_pdf_when_present(tmp_path, monkeypatch):
    dataset, converted = tmp_path / "data", tmp_path / "converted"
    dataset.mkdir()
    monkeypatch.setattr(config, "PAGES_DIR", tmp_path / "pages")
    monkeypatch.setattr(config, "CONVERTED_DIR", converted)
    monkeypatch.setattr(ingest.shutil, "which", lambda name: None)  # no LibreOffice
    document = docx.Document()
    document.add_paragraph("Borehole log without inline images")
    document.save(dataset / "report.docx")
    doc = ingest.SourceDoc("report", "report.docx", "docx")

    assert ingest.render(doc, dataset) == []  # fallback: no embedded images

    converted.mkdir()
    _pdf(converted / "report.pdf", "Sheet 1 of 2", "Sheet 2 of 2")
    pages = ingest.render(doc, dataset)  # new source → rendered again
    assert [p.page for p in pages] == [1, 2]
    assert "Sheet 2 of 2" in pages[1].text
    assert ingest.render(doc, dataset) == pages  # cached now


def test_pdf_next_to_docx_is_its_own_document(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PAGES_DIR", tmp_path / "pages")
    monkeypatch.setattr(config, "CONVERTED_DIR", tmp_path / "converted")
    monkeypatch.setattr(ingest.shutil, "which", lambda name: None)
    dataset = tmp_path / "data"
    dataset.mkdir()
    docx.Document().save(dataset / "r.docx")
    _pdf(dataset / "r.pdf", "Reference copy")
    docs = {d.path: d for d in ingest.discover(dataset)}
    assert docs["r.pdf"].kind == "pdf" and docs["r.docx"].kind == "docx"
    assert ingest.converted_pdf(docs["r.docx"], dataset) is None
    assert ingest.render(docs["r.docx"], dataset) == []  # not rendered from the reference copy


def test_upload_conversions_stay_out_of_the_dataset_folder(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "CONVERTED_DIR", tmp_path / "converted")
    monkeypatch.setattr(config, "UPLOADS_DIR", tmp_path / "uploads")
    assert ingest.converted_dir(tmp_path / "uploads" / "upload-1") == tmp_path / "converted" / "uploads" / "upload-1"
    assert ingest.converted_dir(tmp_path / "dataset") == tmp_path / "converted"


def test_discover_keeps_ids_unique(tmp_path):
    for name in ("plan.pdf", "plan.jpg"):
        (tmp_path / name).write_bytes(b"x")
    docs = {d.path: d.doc_id for d in ingest.discover(tmp_path)}
    assert sorted(docs.values()) == ["plan", "plan-pdf"]


def test_replaced_conversion_is_rendered_again(tmp_path, monkeypatch):
    import os
    converted = tmp_path / "converted"
    converted.mkdir()
    monkeypatch.setattr(config, "PAGES_DIR", tmp_path / "pages")
    monkeypatch.setattr(config, "CONVERTED_DIR", converted)
    docx.Document().save(tmp_path / "r.docx")
    doc = ingest.SourceDoc("r", "r.docx", "docx")
    _pdf(converted / "r.pdf", "sheet 1")
    assert len(ingest.render(doc, tmp_path)) == 1
    _pdf(converted / "r.pdf", "sheet 1", "sheet 2")  # corrected rendering saved over the old one
    os.utime(converted / "r.pdf", ns=(1, 2_000_000_000_000_000_000))
    assert len(ingest.render(doc, tmp_path)) == 2


def test_text_log_keeps_umlauts_and_wraps_long_lines(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PAGES_DIR", tmp_path / "pages")
    line = "0.00 - 1.20 Lehm, br\u00e4unlich, schluffig \u00b1 " + "Kies " * 40
    (tmp_path / "layers.txt").write_bytes(("B 1\n" + line + "\n").encode("cp1252"))  # not UTF-8
    (doc,) = ingest.discover(tmp_path)
    assert doc.kind == "text"
    pages = ingest.render(doc, tmp_path)
    assert len(pages) == 1
    assert "br\u00e4unlich" in pages[0].text and "\u00b1" in pages[0].text
    assert pages[0].text.count("Kies") == 40  # wrapped, not cut at the page edge


def test_spreadsheet_rows_become_text_lines_per_sheet(tmp_path, monkeypatch):
    import openpyxl

    monkeypatch.setattr(config, "PAGES_DIR", tmp_path / "pages")
    book = openpyxl.Workbook()
    book.active.title = "Layers"
    book.active.append(["B 1", 0.5, 1.2, "Lehm, br\u00e4unlich", None, None])
    book.active.append([])
    book.create_sheet("Position").append(["E", 412000, "N", 5320000])
    book.save(tmp_path / "data.xlsx")
    (doc,) = ingest.discover(tmp_path)
    assert doc.kind == "sheet"
    text = ingest.render(doc, tmp_path)[0].text
    assert "[Layers]\nB 1 | 0.5 | 1.2 | Lehm, br\u00e4unlich\n" in text  # trailing empty cells dropped
    assert "[Position]\nE | 412000 | N | 5320000" in text


@pytest.mark.skipif(not shutil.which("soffice") and not shutil.which("libreoffice"), reason="LibreOffice not installed")
def test_docx_table_is_rendered_through_libreoffice(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "PAGES_DIR", tmp_path / "pages")
    monkeypatch.setattr(config, "CONVERTED_DIR", tmp_path / "converted")
    dataset = tmp_path / "data"
    dataset.mkdir()
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    for (row, col), text in {(0, 0): "0.00-1.20", (0, 1): "Fill", (1, 0): "1.20-4.50", (1, 1): "Clay"}.items():
        table.cell(row, col).text = text
    document.save(dataset / "log.docx")
    pages = ingest.render(ingest.SourceDoc("log", "log.docx", "docx"), dataset)
    assert len(pages) == 1 and "1.20-4.50" in pages[0].text and "Clay" in pages[0].text
